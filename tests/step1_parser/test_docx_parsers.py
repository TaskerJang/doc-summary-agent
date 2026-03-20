"""
DOCX 파서 비교 테스트
대상 라이브러리: python-docx / docx2python

.doc 입력 시 LibreOffice CLI로 .docx 변환 후 처리 (REQ-02)
⚠️  LLM·VLM 사용 금지
"""
import subprocess
import time
import traceback
from pathlib import Path
from config import DOCS, output_path


def convert_doc_to_docx(doc_path: Path) -> Path:
    """LibreOffice CLI로 .doc → .docx 변환"""
    out_dir = doc_path.parent
    result = subprocess.run(
        ["libreoffice", "--headless", "--convert-to", "docx",
         "--outdir", str(out_dir), str(doc_path)],
        capture_output=True, text=True
    )
    if result.returncode != 0:
        raise RuntimeError(f"LibreOffice 변환 실패: {result.stderr}")
    return doc_path.with_suffix(".docx")


# ── 1. python-docx ────────────────────────────────────────────
def parse_python_docx(docx_path: Path) -> str:
    from docx import Document
    doc = Document(docx_path)
    lines = []
    for para in doc.paragraphs:
        if para.text.strip():
            prefix = "## " if para.style.name.startswith("Heading") else ""
            lines.append(f"{prefix}{para.text}")
    for table in doc.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(lines)


# ── 2. docx2python ────────────────────────────────────────────
def parse_docx2python(docx_path: Path) -> str:
    from docx2python import docx2python
    result = docx2python(docx_path)
    lines = []
    for section in result.body:
        for paragraph in section:
            for run in paragraph:
                if run.strip():
                    lines.append(run)
    return "\n".join(lines)


# ── 실행 ─────────────────────────────────────────────────────
def run():
    parsers = {
        "python-docx":  parse_python_docx,
        "docx2python":  parse_docx2python,
    }

    doc_targets = {
        "doc_fss": DOCS["doc_fss"],
    }

    for doc_key, doc_path in doc_targets.items():
        if not doc_path.exists():
            print(f"[SKIP] 문서 없음: {doc_path.name}")
            continue

        # .doc → .docx 변환
        if doc_path.suffix.lower() == ".doc":
            print(f"🔄 LibreOffice 변환 중: {doc_path.name}")
            try:
                doc_path = convert_doc_to_docx(doc_path)
                print(f"   → {doc_path.name}")
            except Exception as e:
                print(f"❌ 변환 실패: {e}")
                continue

        print(f"\n{'='*60}")
        print(f"📄 {doc_path.name}")
        print(f"{'='*60}")

        for lib_name, parse_fn in parsers.items():
            print(f"  ▶ {lib_name} ... ", end="", flush=True)
            try:
                t0 = time.perf_counter()
                result = parse_fn(doc_path)
                elapsed = time.perf_counter() - t0

                out = output_path(lib_name, doc_key)
                out.write_text(result, encoding="utf-8")
                print(f"✅ {elapsed:.2f}s | {len(result):,}자 → {out.name}")

            except Exception as e:
                print(f"❌ 실패: {e}")
                traceback.print_exc()


if __name__ == "__main__":
    run()
