"""
HWP 파서 비교 테스트
대상:
  A. pyhwp (hwp5html)  ← 최종 선정 (순수 Python, LibreOffice 불필요)
  B. LibreOffice CLI   ← 실패 기록용
  C. libhwp            ← Rust 패닉 기록용
  D. python-hwpx       ← .hwp 미지원 기록용

⚠️  LLM·VLM 사용 금지
"""
import html
import re
import shutil
import subprocess
import sys
import tempfile
import time
import traceback
from pathlib import Path
from config import DOCS, output_path

if sys.platform == "win32":
    LIBREOFFICE_BIN = r"C:\Program Files\LibreOffice\program\soffice.exe"
else:
    LIBREOFFICE_BIN = "libreoffice"

_HEADING_PREFIX = {
    "Heading 1": "# ",
    "Heading 2": "## ",
    "Heading 3": "### ",
}


# ── A. pyhwp (hwp5html) ───────────────────────────────────────
def parse_pyhwp(hwp_path: Path) -> str:
    """hwp5html로 임시 디렉토리에 xhtml 출력 후 텍스트·표 파싱"""
    tmp_dir = Path(tempfile.mkdtemp())
    try:
        result = subprocess.run(
            ["hwp5html", "--output", str(tmp_dir), str(hwp_path)],
            capture_output=True,
        )
        if result.returncode != 0:
            stderr = result.stderr.decode("utf-8", errors="replace")
            raise RuntimeError(f"hwp5html 실패: {stderr}")

        html_files = sorted(tmp_dir.rglob("*.xhtml")) or sorted(tmp_dir.rglob("*.html"))
        if not html_files:
            raise FileNotFoundError("hwp5html 출력 파일 없음")

        sections = []
        for html_file in html_files:
            raw = html_file.read_text(encoding="utf-8", errors="replace")
            text = re.sub(r"<[^>]+>", " ", raw)
            text = html.unescape(text)
            text = re.sub(r"\n{3,}", "\n\n", text)
            text = re.sub(r" {2,}", " ", text)
            sections.append(text.strip())

        return "\n\n".join(sections)
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)


# ── B. LibreOffice CLI (실패 기록용) ──────────────────────────
def parse_libreoffice(hwp_path: Path) -> str:
    from docx import Document

    out_dir = hwp_path.parent
    result = subprocess.run(
        [LIBREOFFICE_BIN, "--headless", "--convert-to", "docx",
         "--outdir", str(out_dir), str(hwp_path)],
        capture_output=True, text=True,
    )
    if result.returncode != 0:
        raise RuntimeError(f"LibreOffice 변환 실패: {result.stderr}")
    docx_path = hwp_path.with_suffix(".docx")
    if not docx_path.exists():
        raise FileNotFoundError(f"변환 결과 파일 없음: {docx_path}")
    try:
        doc = Document(docx_path)
        lines = []
        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            prefix = _HEADING_PREFIX.get(para.style.name, "")
            lines.append(f"{prefix}{text}")
        for table in doc.tables:
            for row in table.rows:
                lines.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(lines)
    finally:
        if docx_path.exists():
            docx_path.unlink()


# ── C. libhwp (Rust 패닉 기록용) ─────────────────────────────
def parse_libhwp(hwp_path: Path) -> str:
    try:
        from libhwp import HWPReader
        reader = HWPReader(str(hwp_path))
        lines = [str(p).strip() for p in reader.get_paragraphs() if str(p).strip()]
        return "\n".join(lines)
    except ImportError:
        return "[libhwp] 미설치"
    except BaseException as e:
        return f"[libhwp] Rust 패닉: {e}"


# ── D. python-hwpx (.hwp 미지원 기록용) ──────────────────────
def parse_python_hwpx(hwp_path: Path) -> str:
    if hwp_path.suffix.lower() == ".hwp":
        return "[python-hwpx] .hwp 미지원 — .hwpx 전용 라이브러리"
    try:
        import hwpx
        return hwpx.load(str(hwp_path)).get_text()
    except ImportError:
        return "[python-hwpx] 미설치"


# ── 실행 ──────────────────────────────────────────────────────
def run():
    parsers = {
        "pyhwp":       parse_pyhwp,        # 최종 선정
        "libreoffice": parse_libreoffice,   # 실패 기록용
        "libhwp":      parse_libhwp,        # 실패 기록용
        "python-hwpx": parse_python_hwpx,   # 실패 기록용
    }

    hwp_targets = {
        "hwp_nonghyup": DOCS["hwp_nonghyup"],
    }

    for doc_key, doc_path in hwp_targets.items():
        if not doc_path.exists():
            print(f"[SKIP] 문서 없음: {doc_path.name}")
            continue

        print(f"\n{'='*60}")
        print(f"📄 {doc_path.name}")
        print(f"{'='*60}")

        for lib_name, parse_fn in parsers.items():
            print(f"  ▶ {lib_name} ... ", end="", flush=True)
            try:
                t0 = time.perf_counter()
                result = parse_fn(doc_path)
                elapsed = time.perf_counter() - t0

                out = output_path(lib_name, doc_key)
                out.write_text(result, encoding="utf-8")
                print(f"✅ {elapsed:.2f}s | {len(result):,}자 → {out.name}")

            except BaseException as e:
                print(f"❌ 실패: {e}")
                traceback.print_exc()


if __name__ == "__main__":
    run()